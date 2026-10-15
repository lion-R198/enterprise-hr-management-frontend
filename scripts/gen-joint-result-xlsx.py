# -*- coding: utf-8 -*-
"""
生成《前后端联调测试结果.xlsx》

数据来源：verify/joint-result.json（联调测试台真实运行结果 + 服务端代理日志交叉验证）

交付形态严格遵循需求：**只列两列**
    ① 前端请求       —— 该请求的方法与接口路径
    ② 是否调用后端接口 —— 是 / 否（双证据判定：客户端网络层探针 ∩ 服务端代理日志）
其余信息（状态码、业务报文、所属模块、执行角色等）一律不落表。

用法：python scripts/gen-joint-result-xlsx.py
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "verify", "joint-result.json")
OUT = os.path.join(ROOT, "前后端联调测试结果.xlsx")

HEADERS = ["前端请求", "是否调用后端接口"]

# ---------------- 读取联调结果 ----------------
if not os.path.exists(SRC):
    sys.exit(f"找不到联调结果文件：{SRC}")

with open(SRC, encoding="utf-8") as f:
    data = json.load(f)

cases = data.get("cases") or []
if not cases:
    sys.exit("联调结果中没有用例记录")

# ---------------- 组装两列数据 ----------------
rows = []
for c in cases:
    path = str(c.get("path") or "")
    full = path if path.startswith("/api") else "/api" + path
    rows.append([f"{str(c.get('method') or '').upper()} {full}",
                 "是" if c.get("reachedBackend") else "否"])

# ---------------- 写入工作簿 ----------------
wb = Workbook()
ws = wb.active
ws.title = "联调测试结果"

thin = Side(style="thin", color="B7C4D6")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

ws.append(HEADERS)
for row in rows:
    ws.append(row)

# 表头样式
head_fill = PatternFill("solid", fgColor="2F5597")
for idx in range(1, len(HEADERS) + 1):
    cell = ws.cell(row=1, column=idx)
    cell.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    cell.fill = head_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border
ws.row_dimensions[1].height = 26

# 数据区样式
band = PatternFill("solid", fgColor="F2F6FC")
for r in range(2, len(rows) + 2):
    for cidx in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=r, column=cidx)
        cell.font = Font(name="微软雅黑", size=10)
        cell.border = border
        if cidx == 1:
            cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        else:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        if r % 2 == 0:
            cell.fill = band
    ws.row_dimensions[r].height = 20

ws.column_dimensions[get_column_letter(1)].width = 52
ws.column_dimensions[get_column_letter(2)].width = 22

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(rows) + 1}"
ws.sheet_view.showGridLines = False

wb.save(OUT)

yes = sum(1 for r in rows if r[1] == "是")
print(f"已生成：{OUT}")
print(f"共 {len(rows)} 条前端请求｜调用后端接口 {yes} 条｜未调用后端接口 {len(rows) - yes} 条")
